"""Create a reproducible beauty-and-wellness EDA and feature-engineering report.

The default dataset is synthetic because no real preprocessed data was supplied.
Every report table, figure, statistical result, and model metric is calculated
from the generated data at runtime.
"""

from __future__ import annotations

from pathlib import Path

import matplotlib

matplotlib.use("Agg")

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import seaborn as sns
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from scipy.stats import chi2_contingency, kruskal, mannwhitneyu
from sklearn.linear_model import LogisticRegression
from sklearn.metrics import (
    accuracy_score,
    brier_score_loss,
    roc_auc_score,
)
from sklearn.model_selection import train_test_split
from sklearn.pipeline import make_pipeline
from sklearn.preprocessing import StandardScaler


SEED = 2026
AS_OF_DATE = pd.Timestamp("2025-12-31")
OUTPUT_DIR = Path(__file__).resolve().parent / "output"
N_CUSTOMERS = 1_500


def sigmoid(values: np.ndarray) -> np.ndarray:
    """Numerically stable enough for the bounded synthetic inputs used here."""
    return 1.0 / (1.0 + np.exp(-values))


def build_synthetic_data(
    n_customers: int = N_CUSTOMERS, seed: int = SEED
) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Generate historical order lines and a customer-level repeat-purchase label."""
    rng = np.random.default_rng(seed)
    categories = np.array(
        ["Skincare", "Haircare", "Makeup", "Supplements", "Spa Services"]
    )
    channels = np.array(["Online", "Salon", "Marketplace", "Retail"])
    regions = np.array(["Northeast", "South", "Midwest", "West"])
    order_rows: list[dict[str, object]] = []
    customer_rows: list[dict[str, object]] = []

    for customer_number in range(n_customers):
        customer_id = f"C{customer_number + 1:05d}"
        age = int(rng.integers(18, 71))
        channel = str(rng.choice(channels, p=[0.48, 0.20, 0.19, 0.13]))
        region = str(rng.choice(regions))
        loyalty_member = int(rng.random() < 0.34)
        engagement = float(rng.normal())
        tenure_days = int(rng.integers(180, 1_461))
        order_count = int(
            np.clip(rng.poisson(2.7 + 0.65 * engagement + 0.6 * loyalty_member), 1, 12)
        )

        order_dates = AS_OF_DATE - pd.to_timedelta(
            rng.integers(0, tenure_days + 1, size=order_count), unit="D"
        )
        preferred_category = int(rng.integers(0, len(categories)))
        category_probabilities = np.full(len(categories), 0.12, dtype=float)
        category_probabilities[preferred_category] = 0.52
        category_probabilities /= category_probabilities.sum()
        customer_orders: list[dict[str, object]] = []

        for order_number, order_date in enumerate(order_dates):
            category = str(rng.choice(categories, p=category_probabilities))
            discount = float(
                np.clip(
                    rng.beta(1.7, 5.0) + (0.09 if channel == "Marketplace" else 0.0),
                    0.0,
                    0.65,
                )
            )
            category_price = {
                "Skincare": 42,
                "Haircare": 34,
                "Makeup": 31,
                "Supplements": 38,
                "Spa Services": 86,
            }[category]
            gross_amount = float(
                rng.lognormal(np.log(category_price), 0.42)
                * (1 + 0.11 * loyalty_member)
            )
            amount = round(gross_amount * (1 - discount), 2)
            satisfaction = float(
                np.clip(
                    4.05
                    + 0.22 * engagement
                    + 0.12 * loyalty_member
                    - 0.65 * discount
                    + rng.normal(0, 0.48),
                    1.0,
                    5.0,
                )
            )
            row = {
                "order_id": f"{customer_id}-O{order_number + 1:02d}",
                "customer_id": customer_id,
                "order_date": pd.Timestamp(order_date),
                "category": category,
                "channel": channel,
                "region": region,
                "order_amount": amount,
                "discount_pct": round(discount, 4),
                "rating": round(satisfaction, 2),
                "loyalty_member": loyalty_member,
            }
            order_rows.append(row)
            customer_orders.append(row)

        customer_order_frame = pd.DataFrame(customer_orders)
        days_since_last_order = int(
            (AS_OF_DATE - customer_order_frame["order_date"].max()).days
        )
        category_diversity = int(customer_order_frame["category"].nunique())
        wellness_share = float(
            customer_order_frame["category"].isin(["Supplements", "Spa Services"]).mean()
        )
        discount_order_share = float(
            (customer_order_frame["discount_pct"] >= 0.15).mean()
        )
        average_rating = float(customer_order_frame["rating"].mean())
        total_spend = float(customer_order_frame["order_amount"].sum())
        average_order_value = float(customer_order_frame["order_amount"].mean())
        orders_per_year = float(order_count / (tenure_days / 365.25))

        log_odds = (
            -0.72
            + 0.19 * (order_count - 3)
            - 0.010 * (days_since_last_order - 60)
            + 0.70 * (average_rating - 4.0)
            + 0.42 * (category_diversity - 2.0)
            + 0.62 * (wellness_share - 0.30)
            - 0.48 * (discount_order_share - 0.30)
            + 0.25 * loyalty_member
        )
        repeat_probability = float(sigmoid(np.array([log_odds]))[0])
        repeat_purchase = int(rng.random() < repeat_probability)
        customer_rows.append(
            {
                "customer_id": customer_id,
                "age": age,
                "channel": channel,
                "region": region,
                "loyalty_member": loyalty_member,
                "tenure_days": tenure_days,
                "total_orders": order_count,
                "total_spend": round(total_spend, 2),
                "avg_order_value": round(average_order_value, 2),
                "recency_days": days_since_last_order,
                "category_diversity": category_diversity,
                "wellness_share": wellness_share,
                "discount_order_share": discount_order_share,
                "avg_rating": average_rating,
                "orders_per_year": orders_per_year,
                "repeat_purchase_90d": repeat_purchase,
            }
        )

    transactions = pd.DataFrame(order_rows).sort_values(
        ["order_date", "order_id"], ignore_index=True
    )
    customers = pd.DataFrame(customer_rows)
    return transactions, customers


def make_visualizations(
    transactions: pd.DataFrame, customers: pd.DataFrame, output_dir: Path
) -> dict[str, Path]:
    """Save the principal visuals and return their paths."""
    sns.set_theme(style="whitegrid", palette="Set2")
    figure_paths = {
        "monthly": output_dir / "monthly_revenue.png",
        "category": output_dir / "category_revenue.png",
        "correlation": output_dir / "feature_correlation.png",
    }

    monthly = (
        transactions.assign(month=transactions["order_date"].dt.to_period("M").dt.to_timestamp())
        .groupby("month", as_index=False)["order_amount"]
        .sum()
    )
    fig, ax = plt.subplots(figsize=(9.0, 4.2))
    sns.lineplot(data=monthly, x="month", y="order_amount", marker="o", ax=ax)
    ax.set(title="Synthetic monthly revenue", xlabel="Order month", ylabel="Revenue ($)")
    fig.autofmt_xdate()
    fig.tight_layout()
    fig.savefig(figure_paths["monthly"], dpi=160)
    plt.close(fig)

    category_revenue = (
        transactions.groupby("category", as_index=False)["order_amount"]
        .sum()
        .sort_values("order_amount", ascending=False)
    )
    fig, ax = plt.subplots(figsize=(8.0, 4.2))
    sns.barplot(data=category_revenue, x="order_amount", y="category", ax=ax)
    ax.set(title="Synthetic revenue by category", xlabel="Revenue ($)", ylabel="")
    fig.tight_layout()
    fig.savefig(figure_paths["category"], dpi=160)
    plt.close(fig)

    correlation_columns = [
        "total_orders",
        "total_spend",
        "recency_days",
        "category_diversity",
        "wellness_share",
        "discount_order_share",
        "avg_rating",
        "orders_per_year",
        "repeat_purchase_90d",
    ]
    fig, ax = plt.subplots(figsize=(8.2, 6.4))
    sns.heatmap(
        customers[correlation_columns].corr(),
        cmap="vlag",
        center=0,
        annot=True,
        fmt=".2f",
        square=True,
        ax=ax,
        annot_kws={"size": 7},
    )
    ax.set_title("Customer feature correlations (Pearson)")
    fig.tight_layout()
    fig.savefig(figure_paths["correlation"], dpi=160)
    plt.close(fig)
    return figure_paths


def compare_models(customers: pd.DataFrame) -> tuple[pd.DataFrame, float]:
    """Compare held-out logistic models with conventional and engineered features."""
    baseline_features = [
        "age",
        "tenure_days",
        "total_orders",
        "total_spend",
        "loyalty_member",
    ]
    engineered_features = baseline_features + [
        "recency_days",
        "avg_order_value",
        "category_diversity",
        "wellness_share",
        "discount_order_share",
        "avg_rating",
        "orders_per_year",
    ]
    target = customers["repeat_purchase_90d"]
    indices = np.arange(len(customers))
    train_indices, test_indices = train_test_split(
        indices, test_size=0.25, random_state=SEED, stratify=target
    )
    rows: list[dict[str, float | str]] = []
    for name, features in (
        ("Baseline", baseline_features),
        ("Engineered", engineered_features),
    ):
        model = make_pipeline(
            StandardScaler(),
            LogisticRegression(max_iter=2_000, random_state=SEED),
        )
        model.fit(customers.iloc[train_indices][features], target.iloc[train_indices])
        probabilities = model.predict_proba(customers.iloc[test_indices][features])[:, 1]
        predictions = (probabilities >= 0.5).astype(int)
        rows.append(
            {
                "model": name,
                "test_roc_auc": roc_auc_score(target.iloc[test_indices], probabilities),
                "test_accuracy": accuracy_score(target.iloc[test_indices], predictions),
                "test_brier_score": brier_score_loss(
                    target.iloc[test_indices], probabilities
                ),
            }
        )
    metrics = pd.DataFrame(rows)
    auc_gain = float(
        metrics.loc[metrics["model"] == "Engineered", "test_roc_auc"].iloc[0]
        - metrics.loc[metrics["model"] == "Baseline", "test_roc_auc"].iloc[0]
    )
    return metrics, auc_gain


def format_p_value(value: float) -> str:
    return "<0.001" if value < 0.001 else f"{value:.3f}"


def add_table(document: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Light Shading Accent 1"
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text
    for row in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = str(text)


def add_bullet(document: Document, text: str) -> None:
    document.add_paragraph(text, style="List Bullet")


def build_report(
    transactions: pd.DataFrame,
    customers: pd.DataFrame,
    metrics: pd.DataFrame,
    auc_gain: float,
    figure_paths: dict[str, Path],
    output_path: Path,
) -> None:
    """Create the Word report using the computed synthetic-data results."""
    document = Document()
    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)
    styles["Normal"].font.color.rgb = RGBColor(38, 48, 61)
    for style_name in ("Heading 1", "Heading 2"):
        styles[style_name].font.name = "Arial"
        styles[style_name].font.color.rgb = RGBColor(31, 78, 121)

    title = document.add_heading(
        "Week 4: Exploratory Data Analysis and Feature Engineering", 0
    )
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph(
        "Beauty & wellness customer behavior | Reproducible worked example"
    )
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph(
        "Data provenance: simulated, not observed. No source dataset was available "
        "in the workspace. The figures and numerical results below are generated "
        "from the seeded synthetic data produced by the companion Python script."
    )
    document.add_heading("1. Executive summary", level=1)
    repeat_rate = customers["repeat_purchase_90d"].mean()
    revenue = transactions["order_amount"].sum()
    category_revenue = transactions.groupby("category")["order_amount"].sum()
    leading_category = str(category_revenue.idxmax())
    leading_category_share = float(category_revenue.max() / revenue)
    engineered_auc = float(
        metrics.loc[metrics["model"] == "Engineered", "test_roc_auc"].iloc[0]
    )
    baseline_auc = float(
        metrics.loc[metrics["model"] == "Baseline", "test_roc_auc"].iloc[0]
    )
    document.add_paragraph(
        f"The reproducible example contains {len(customers):,} simulated customers "
        f"and {len(transactions):,} historical orders dated through "
        f"{AS_OF_DATE:%Y-%m-%d}. Simulated total revenue is ${revenue:,.2f}; "
        f"{leading_category} contributes {leading_category_share:.1%} of it. "
        f"The generated 90-day repeat-purchase rate is {repeat_rate:.1%}. "
        f"On a stratified 25% holdout, engineered customer behavior features "
        f"produce ROC-AUC {engineered_auc:.3f} versus {baseline_auc:.3f} for the "
        f"baseline (difference {auc_gain:+.3f}). This is a demonstration of the "
        "workflow, not evidence of production uplift."
    )

    document.add_heading("2. Business question and analytical scope", level=1)
    document.add_paragraph(
        "Primary question: which observable purchase-history patterns are useful "
        "for identifying customers likely to place another order within 90 days? "
        "Secondary questions: how do category mix, discount use, satisfaction, "
        "and recency vary across customer groups, and which patterns suggest "
        "testable retention or cross-category opportunities?"
    )
    add_bullet(
        document,
        "Unit of analysis: order line for revenue and category analysis; customer "
        "for RFM summaries and the repeat-purchase model.",
    )
    add_bullet(
        document,
        f"Observation cutoff: {AS_OF_DATE:%Y-%m-%d}. Historical behavior is "
        "aggregated at that cutoff; the label is a simulated next-90-day outcome.",
    )
    add_bullet(
        document,
        "Population: customers with at least one historical order in the "
        "synthetic extract; therefore the example does not estimate acquisition "
        "or first-purchase conversion.",
    )
    add_bullet(
        document,
        "Business use: prioritize experiments and customer-service outreach; "
        "do not use age, region, or inferred wellness preferences for sensitive "
        "decisions without legal, fairness, and privacy review.",
    )

    document.add_heading("3. Data audit and quality checks", level=1)
    missing_order_cells = int(transactions.isna().sum().sum())
    missing_customer_cells = int(customers.isna().sum().sum())
    duplicate_order_ids = int(transactions["order_id"].duplicated().sum())
    invalid_dates = int(
        (
            (transactions["order_date"] > AS_OF_DATE)
            | (transactions["order_date"] < AS_OF_DATE - pd.Timedelta(days=1_460))
        ).sum()
    )
    audit_rows = [
        ["Order rows", f"{len(transactions):,}"],
        ["Customer rows", f"{len(customers):,}"],
        ["Missing cells: orders / customers", f"{missing_order_cells} / {missing_customer_cells}"],
        ["Duplicate order IDs", str(duplicate_order_ids)],
        ["Dates outside configured history window", str(invalid_dates)],
        ["Observed order date range", f"{transactions['order_date'].min():%Y-%m-%d} to {transactions['order_date'].max():%Y-%m-%d}"],
    ]
    document.add_paragraph(
        "Run structural checks before interpreting charts: schema and data types, "
        "unique order keys, customer-key referential integrity, missingness, "
        "duplicate lines, date range, nonnegative spend, discount bounds, and "
        "rating bounds. For real exports, reconcile sales totals with finance "
        "and confirm whether refunds, taxes, returns, and canceled orders are "
        "included."
    )
    add_table(document, ["Audit item", "Observed result"], audit_rows)

    document.add_heading("4. EDA workflow and statistical methods", level=1)
    document.add_heading("4.1 Univariate and distributional review", level=2)
    document.add_paragraph(
        "Report count, mean, standard deviation, median, IQR, minimum, maximum, "
        "and missingness for spend, discount, rating, recency, and order count. "
        "Use histograms/KDEs for shape and boxplots for skew and extreme values. "
        "Revenue is typically right-skewed, so retain the raw scale for monetary "
        "interpretation and use log1p only for suitable model inputs. Flag "
        "possible outliers with IQR fences, then verify against business rules "
        "rather than deleting automatically."
    )
    document.add_heading("4.2 Segmentation and relationships", level=2)
    document.add_paragraph(
        "Compare customer spend, frequency, wellness share, and repeat rate by "
        "acquisition channel, loyalty status, region, and category. Use median "
        "and IQR alongside means. Use Spearman correlation for monotonic "
        "relationships between skewed numeric measures; Pearson heatmaps are "
        "included here as a screening visualization only. Inspect pair plots "
        "on a stratified sample when interactions are suspected."
    )
    document.add_heading("4.3 Tests and uncertainty", level=2)
    document.add_paragraph(
        "For two independent groups (for example, repeat versus non-repeat "
        "customer spend), use Mann–Whitney U when distributional assumptions "
        "for a t-test are unsuitable; report group medians and an effect size, "
        "not only a p-value. For three or more groups (such as order count by "
        "channel), use Kruskal–Wallis, followed by multiplicity-adjusted "
        "pairwise tests if significant. For category-by-channel counts, use "
        "Pearson chi-square and Cramér's V; combine sparse levels or use an "
        "exact/Monte Carlo method if expected cell counts are too small. "
        "Predefine comparisons, report confidence intervals, and distinguish "
        "exploration from confirmatory inference."
    )
    document.add_heading("4.4 Worked statistical examples", level=2)
    repeat_spend = customers.loc[
        customers["repeat_purchase_90d"] == 1, "total_spend"
    ]
    nonrepeat_spend = customers.loc[
        customers["repeat_purchase_90d"] == 0, "total_spend"
    ]
    spend_test = mannwhitneyu(repeat_spend, nonrepeat_spend, alternative="two-sided")
    rank_biserial = 2 * spend_test.statistic / (
        len(repeat_spend) * len(nonrepeat_spend)
    ) - 1
    channel_groups = [
        customers.loc[customers["channel"] == channel, "total_orders"]
        for channel in sorted(customers["channel"].unique())
    ]
    channel_test = kruskal(*channel_groups)
    contingency = pd.crosstab(transactions["channel"], transactions["category"])
    chi_square, chi_p, _, expected = chi2_contingency(contingency)
    cramers_v = float(
        np.sqrt(
            chi_square
            / (contingency.to_numpy().sum() * min(contingency.shape[0] - 1, contingency.shape[1] - 1))
        )
    )
    stats_rows = [
        [
            "Customer spend: repeat vs. non-repeat (Mann–Whitney U)",
            f"${repeat_spend.median():,.2f} vs. ${nonrepeat_spend.median():,.2f}",
            f"U={spend_test.statistic:,.0f}; p={format_p_value(float(spend_test.pvalue))}; rank-biserial r={rank_biserial:.3f}",
        ],
        [
            "Order counts across channels (Kruskal–Wallis)",
            "Independent customer groups",
            f"H={channel_test.statistic:.2f}; p={format_p_value(float(channel_test.pvalue))}",
        ],
        [
            "Order category by channel (chi-square)",
            f"Minimum expected cell={expected.min():.1f}",
            f"chi-square={chi_square:.2f}; p={format_p_value(float(chi_p))}; Cramér's V={cramers_v:.3f}",
        ],
    ]
    add_table(document, ["Worked comparison", "Observed groups/check", "Test result"], stats_rows)
    document.add_paragraph(
        "These tests demonstrate the calculation on seeded simulated data. "
        "They do not establish causal effects, and statistical significance "
        "does not by itself establish a commercially meaningful difference."
    )

    document.add_heading("5. Visualization plan and worked results", level=1)
    document.add_paragraph(
        "Use pandas for schema checks, aggregation, and reproducible tables; "
        "Matplotlib for explicit axes, labels, and report-ready figures; Seaborn "
        "for distribution, categorical, and correlation views; SciPy for "
        "nonparametric and contingency-table tests; and scikit-learn for "
        "preprocessing pipelines and held-out evaluation."
    )
    for path, caption in (
        (figure_paths["monthly"], "Figure 1. Monthly synthetic revenue. Look for seasonality and abrupt changes; a synthetic curve is not a forecast."),
        (figure_paths["category"], "Figure 2. Synthetic revenue by product/service category. Validate returns, price mix, and exposure before interpreting category differences."),
        (figure_paths["correlation"], "Figure 3. Pearson correlations among customer-level measures. Correlation is descriptive, not causal; correlated features may require regularization."),
    ):
        document.add_picture(str(path), width=Inches(6.15))
        document.add_paragraph(caption, style="Caption")

    document.add_heading("6. Feature-engineering plan", level=1)
    add_table(
        document,
        ["Feature", "Construction", "Business signal / caution"],
        [
            ["Recency", "Cutoff date − last order date (days)", "Lapse risk and replenishment timing; make cutoff explicit."],
            ["Frequency", "Distinct orders in trailing 365 days", "Engagement cadence; deduplicate order lines."],
            ["Monetary / AOV", "Net spend sum / distinct orders", "Customer value; exclude tax/refunds consistently."],
            ["Category diversity", "Distinct categories per customer", "Breadth of routine; not proof of cross-sell intent."],
            ["Wellness share", "Supplements and spa order count / all orders", "Wellness affinity proxy; validate taxonomy and consent."],
            ["Discount-order share", "Orders with discount >= 15% / all orders", "Promotion reliance proxy; threshold must be tested."],
            ["Satisfaction", "Mean post-order rating", "Experience signal; address missing/non-random survey response."],
            ["Order cadence", "Orders / observed tenure in years", "Normalizes frequency for customer exposure time."],
            ["Interaction", "Recency × wellness share; discount share × channel", "Tests whether recency or promotion response differs by segment."],
        ],
    )
    document.add_heading("6.1 Concrete transformations", level=2)
    document.add_paragraph(
        "Example for a customer with 4 orders, $240 spend, 2 wellness orders, "
        "and 1 discounted order (at the defined 15% cutoff): AOV = $240 / 4 = "
        "$60; wellness share = 2 / 4 = 0.50; discount-order share = 1 / 4 = "
        "0.25. If the cutoff date is 2025-12-31 and the last order was "
        "2025-11-16, recency = 45 days. These calculations are examples of "
        "feature definitions, not a claim about any real customer."
    )
    document.add_paragraph(
        "For modeling: standardize continuous variables within a training-only "
        "Pipeline; use log1p for nonnegative long-tailed monetary/frequency "
        "variables after checking zeros; one-hot encode nominal channel and "
        "region; keep ordered categories ordinal only when the order has a "
        "defensible meaning. Fit imputers, scalers, encoders, and any feature "
        "selection on training folds only. Avoid target leakage by computing "
        "every feature strictly before the label window."
    )

    document.add_heading("7. Model comparison and business interpretation", level=1)
    model_rows = [
        [
            str(row["model"]),
            f"{row['test_roc_auc']:.3f}",
            f"{row['test_accuracy']:.3f}",
            f"{row['test_brier_score']:.3f}",
        ]
        for _, row in metrics.iterrows()
    ]
    add_table(
        document,
        ["Model", "ROC-AUC (higher)", "Accuracy at 0.5", "Brier score (lower)"],
        model_rows,
    )
    document.add_paragraph(
        f"The engineered model's held-out ROC-AUC change is {auc_gain:+.3f}. "
        "The outcome was generated partly from engineered behavior signals, so "
        "this simulation is intentionally suited to illustrating how richer "
        "history can add predictive information. The metric is not a fair "
        "estimate of real-world lift. In a real project, use a temporal "
        "train/validation/test split, tune thresholds against campaign costs, "
        "calibrate probabilities, compare against a simple business rule, and "
        "report confidence intervals and segment-level error rates."
    )
    document.add_heading("Hypothetical decision scenario", level=2)
    document.add_paragraph(
        "A wellness retailer could rank consented customers by predicted "
        "90-day repeat propensity, then run a randomized holdout test of "
        "replenishment reminders. Customers with high wellness share and "
        "increasing recency could receive routine-specific content; customers "
        "whose purchases are discount-heavy could be tested with a value bundle "
        "rather than a blanket coupon. Measure incremental contribution margin, "
        "opt-outs, and complaints—not clicks alone. The model should support "
        "experimentation and resource allocation, not assert why an individual "
        "behaves as they do."
    )

    document.add_heading("8. Limitations and production validation", level=1)
    for item in (
        "All records and relationships are simulated. The generated outcome embeds known feature effects; therefore model performance is illustrative and potentially optimistic.",
        "The sample omits returns, inventory, product-level ingredients, subscriptions, promotions over time, consent status, and real customer acquisition costs.",
        "The repeat label is generated probabilistically from history, not observed from a future transaction system; real labels require a mature 90-day outcome window.",
        "Customers with no recorded orders are excluded. Production use needs explicit cohort definitions and checks for data coverage, survivorship, and channel attribution.",
        "Before deployment, validate on an out-of-time sample, check calibration and subgroup performance, monitor drift, review privacy and fairness, and test interventions with randomized controls.",
    ):
        add_bullet(document, item)

    document.add_heading("9. Proposed 32-hour work plan", level=1)
    add_table(
        document,
        ["Work block", "Hours", "Output"],
        [
            ["Data understanding, source definitions, and quality audit", "4", "Data dictionary, key/date/duplicate/missingness checks"],
            ["Univariate and customer-segment EDA", "6", "Summary tables and distribution plots"],
            ["Relationships, statistical tests, and uncertainty", "5", "Segment comparisons with effect sizes and caveats"],
            ["Feature definitions and leakage review", "5", "RFM, routine, wellness, discount, satisfaction features"],
            ["Reproducible transformations and model comparison", "6", "Pipeline, temporal validation plan, baseline metrics"],
            ["Business interpretation, report, and QA", "6", "DOCX, charts, reproducibility and limitation review"],
            ["Total", "32", "Within the requested 30–35 hour workload"],
        ],
    )

    document.add_heading("10. Reproduction", level=1)
    document.add_paragraph(
        "Install the packages in requirements.txt and run "
        "`python eda_feature_engineering.py`. The random seed is 2026. The "
        "script writes this report, CSVs, charts, and held-out model metrics "
        "to the output directory. Replace the synthetic generator with the "
        "authorized preprocessed dataset before presenting results as empirical."
    )
    document.save(output_path)


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    transactions, customers = build_synthetic_data()
    if transactions.empty or customers.empty:
        raise RuntimeError("Synthetic data generation returned no records.")
    if transactions["order_id"].duplicated().any():
        raise RuntimeError("Generated order IDs are not unique.")
    if transactions["order_amount"].isna().any() or customers.isna().any().any():
        raise RuntimeError("Unexpected missing values found in generated data.")
    if not transactions["discount_pct"].between(0, 0.65).all():
        raise RuntimeError("Generated discounts fell outside configured bounds.")

    transactions.to_csv(OUTPUT_DIR / "synthetic_transactions.csv", index=False)
    customers.to_csv(OUTPUT_DIR / "synthetic_customer_features.csv", index=False)
    figure_paths = make_visualizations(transactions, customers, OUTPUT_DIR)
    metrics, auc_gain = compare_models(customers)
    metrics.to_csv(OUTPUT_DIR / "model_comparison.csv", index=False)
    build_report(
        transactions,
        customers,
        metrics,
        auc_gain,
        figure_paths,
        OUTPUT_DIR / "week4_eda_feature_engineering.docx",
    )
    print(f"Customers: {len(customers):,}")
    print(f"Orders: {len(transactions):,}")
    print(f"Repeat purchase rate: {customers['repeat_purchase_90d'].mean():.3f}")
    print(metrics.to_string(index=False, float_format=lambda value: f"{value:.3f}"))
    print(f"ROC-AUC change (engineered - baseline): {auc_gain:+.3f}")
    print(f"Report: {OUTPUT_DIR / 'week4_eda_feature_engineering.docx'}")


if __name__ == "__main__":
    main()
